import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import pandas as pd
import numpy as np
import os
import sys
import time
import pickle
sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(BASE_DIR, '..', '第二问'))
sys.path.insert(0, os.path.join(BASE_DIR, '..', '第三问'))
import q2_main as q2m
import q3_main as q3m

plt.rcParams['font.family'] = 'SimHei'
plt.rcParams['axes.unicode_minus'] = False
Gm4 = q3m.a1['光伏发电预测功率'].values

T = 144
OUT = BASE_DIR
Q2_REF = 19793691.58
ANCHOR_A11 = 2844.1e4
ANCHOR_A122 = 2121.8e4
ANCHOR_LEDGER = 2872.3e4
FOUR_DAYS = ['2025-03-20', '2025-06-21', '2025-09-23', '2025-12-21']

print('=' * 72)
print('第四问：波动电价重算（模型结构零改动，仅价格矩阵替换）')
print('=' * 72)
a4 = q3m.a4
P4_raw = a4.iloc[:, 1:].values.astype(float)
print(f'\n0. 附件4: {P4_raw.shape}，均值 {P4_raw.mean():.4f} '
      f'最小 {P4_raw.min():.4f} 最大 {P4_raw.max():.4f}，'
      f'近零(<0.05) {int((P4_raw < 0.05).sum())} 个，'
      f'尖峰(>1.2) {int((P4_raw > 1.2).sum())} 个')


def P4_of(align):
    return P4_raw if align == 'direct' else np.roll(P4_raw, 1, axis=1)


def q2_run(mode, alpha, beta, P4, collect=False, tag=''):
    return q2m.run_year(mode, q3m.p, q3m.Lm, Gm4, q3m.L_all, q3m.G_all,
                        q3m.dates, q3m.mask_out, alpha=alpha, beta=beta,
                        collect=collect, price_mat=P4, tag=tag)


def q3_run(rev_hours, alpha, beta, P4, collect=False, tag=''):
    return q3m.run_year_q3(rev_hours, alpha, beta, collect=collect,
                           price_mat=P4, tag=tag)


print('\n1. 对齐仲裁（A(1,1) 两种列映射 vs 教程锚点 2844.1 万）:')
cand = {}
for align in ['direct', 'roll']:
    st = q2_run('A', 1.0, 1.0, P4_of(align))
    cand[align] = st['plan_cost'].sum() + st['emerg_cost'].sum()
    print(f'   {align:>7}: {cand[align]/1e4:,.1f} 万（vs 锚点差 '
          f'{(cand[align]-ANCHOR_A11)/1e4:+.1f} 万）', flush=True)
ALIGN = min(cand, key=lambda k: abs(cand[k] - ANCHOR_A11))
P4 = P4_of(ALIGN)
print(f'   → 采用对齐: {ALIGN}（列 j ↔ 时段 '
      f'{"j" if ALIGN == "direct" else "(j+1)%144"}）')

print('\n2. 退化自检（固定电价矩阵 tile(附件1) → 必须逐分复现 Q2/Q3 封版数字）:')
t0 = time.time()
st_tile = q2_run('B', 1.22, 1.00, np.tile(q3m.p, (365, 1)))
tot_tile = st_tile['plan_cost'].sum() + st_tile['emerg_cost'].sum()
ok1 = abs(tot_tile - Q2_REF) < 1.0
print(f'   Q2-B(1.22,1.00): {tot_tile:,.2f} vs 参考 {Q2_REF:,.2f} '
      f'{"✓" if ok1 else "✗"}（{time.time()-t0:.0f}s）')

with open(os.path.join(OUT, '_q3v2_S1218_a1.22.pkl'), 'rb') as f:
    q3_ref = pickle.load(f)['total']
t0 = time.time()
st3_tile = q3_run((12, 18), 1.22, 1.00, np.tile(q3m.p, (365, 1)))
ok2 = abs(st3_tile['total'] - q3_ref) < 1.0
print(f'   Q3-S1218(1.22,1.00): {st3_tile["total"]:,.2f} vs pickle {q3_ref:,.2f} '
      f'{"✓" if ok2 else "✗"}（{time.time()-t0:.0f}s）')
if not (ok1 and ok2):
    print('   ✗ 退化自检失败——换价改造有误，禁止进入 Q4 计算')
    sys.exit(1)

def run_stage_final():
    print('\n3. final: 交付口径在波动电价下重跑:')
    t0 = time.time()
    B = q2_run('B', 1.22, 1.00, P4, collect=True, tag='Q4·Q2-B(1.22,1.00)')
    with open(os.path.join(OUT, '_q4_B122.pkl'), 'wb') as f:
        pickle.dump(B, f)
    print(f'      ({time.time()-t0:.0f}s，已存盘)', flush=True)
    t0 = time.time()
    S = q3_run((12, 18), 1.22, 1.00, P4, collect=True,
               tag='Q4·Q3-S1218(1.22,1.00)')
    with open(os.path.join(OUT, '_q4_S1218.pkl'), 'wb') as f:
        pickle.dump(S, f)
    print(f'      ({time.time()-t0:.0f}s，已存盘)', flush=True)
    fill_4_2(B, P4)
    fill_4_3(S, P4)
    return B, S


def fill_4_2(B, P4):
    print('   填写 result4-2.xlsx（Q2-B 口径）:')
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(OUT, '..', '附件', '附件5',
                                             'result4-2.xlsx'))
    ws = wb[wb.sheetnames[0]]
    for j in range(len(q3m.idx_out)):
        i = q3m.idx_out[j]
        ds = str(q3m.out_dates.iloc[j].date())
        x_p = B['detail'][ds]['x_p']
        roll = np.roll(x_p, -1)
        r = 2 + j
        for k in range(T):
            ws.cell(row=r, column=2 + k, value=round(float(roll[k]), 2))
        ws.cell(row=r, column=146, value=round(float(x_p.sum()), 2))
        ws.cell(row=r, column=147, value=round(float(P4[i] @ x_p), 2))
    ws = wb[wb.sheetnames[1]]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    blocks = ['0:00-4:00', '4:00-8:00', '8:00-12:00', '12:00-16:00',
              '16:00-20:00', '20:00-24:00']
    r = 2
    for j in range(len(q3m.idx_out)):
        ds = str(q3m.out_dates.iloc[j].date())
        d = B['detail'][ds]
        for b in range(6):
            sl = slice(b * 24, (b + 1) * 24)
            ws.cell(row=r, column=1,
                    value=q3m.out_dates.iloc[j] if b == 0 else None)
            ws.cell(row=r, column=2, value=blocks[b])
            ws.cell(row=r, column=3, value=round(float(d['chg'][sl].sum()), 2))
            ws.cell(row=r, column=4, value=round(float(d['dis'][sl].sum()), 2))
            if b == 0:
                ws.cell(row=r, column=5, value='0:00')
                ws.cell(row=r, column=6, value=round(float(d['e0']), 2))
            elif b == 1:
                ws.cell(row=r, column=5, value='24:00')
                ws.cell(row=r, column=6, value=round(float(d['soc'][-1]), 2))
            r += 1
    ws = wb[wb.sheetnames[2]]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    r = 2
    n_seg = 0
    for j in range(len(q3m.idx_out)):
        ds = str(q3m.out_dates.iloc[j].date())
        ivs = q2m.merge_emergency(B['detail'][ds]['emerg'])
        for m_idx, (label, energy) in enumerate(ivs):
            ws.cell(row=r, column=1,
                    value=q3m.out_dates.iloc[j] if m_idx == 0 else None)
            ws.cell(row=r, column=2, value=label)
            ws.cell(row=r, column=3, value=round(energy, 2))
            r += 1
            n_seg += 1
    wb.save(os.path.join(OUT, 'result4-2.xlsx'))
    print(f'   已保存 result4-2.xlsx（紧急购电 {n_seg} 段）')


def fill_4_3(S, P4):
    print('   填写 result4-3.xlsx（Q3-S1218 口径）:')
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(OUT, '..', '附件', '附件5',
                                             'result4-3.xlsx'))
    det = S['detail']
    for si, (sheet, field) in enumerate([('计划购电量', 'xhat'),
                                         ('调整购电量', 'x_final')]):
        ws = wb[wb.sheetnames[si]]
        for j in range(len(q3m.idx_out)):
            i = q3m.idx_out[j]
            ds = str(q3m.out_dates.iloc[j].date())
            xv = det[ds][field]
            roll = np.roll(xv, -1)
            r = 2 + j
            for k in range(T):
                ws.cell(row=r, column=2 + k, value=round(float(roll[k]), 2))
            ws.cell(row=r, column=146, value=round(float(xv.sum()), 2))
            ws.cell(row=r, column=147, value=round(float(P4[i] @ xv), 2))
    ws = wb[wb.sheetnames[2]]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    blocks = ['0:00-4:00', '4:00-8:00', '8:00-12:00', '12:00-16:00',
              '16:00-20:00', '20:00-24:00']
    r = 2
    for j in range(len(q3m.idx_out)):
        ds = str(q3m.out_dates.iloc[j].date())
        d = det[ds]
        for b in range(6):
            sl = slice(b * 24, (b + 1) * 24)
            ws.cell(row=r, column=1,
                    value=q3m.out_dates.iloc[j] if b == 0 else None)
            ws.cell(row=r, column=2, value=blocks[b])
            ws.cell(row=r, column=3, value=round(float(d['chg'][sl].sum()), 2))
            ws.cell(row=r, column=4, value=round(float(d['dis'][sl].sum()), 2))
            if b == 0:
                ws.cell(row=r, column=5, value='0:00')
                ws.cell(row=r, column=6, value=round(float(d['e0']), 2))
            elif b == 1:
                ws.cell(row=r, column=5, value='24:00')
                ws.cell(row=r, column=6, value=round(float(d['soc'][-1]), 2))
            r += 1
    ws = wb[wb.sheetnames[3]]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    r = 2
    n_seg = 0
    for j in range(len(q3m.idx_out)):
        ds = str(q3m.out_dates.iloc[j].date())
        ivs = q3m.merge_emergency(det[ds]['emerg'])
        for m_idx, (label, energy) in enumerate(ivs):
            ws.cell(row=r, column=1,
                    value=q3m.out_dates.iloc[j] if m_idx == 0 else None)
            ws.cell(row=r, column=2, value=label)
            ws.cell(row=r, column=3, value=round(energy, 2))
            r += 1
            n_seg += 1
    wb.save(os.path.join(OUT, 'result4-3.xlsx'))
    print(f'   已保存 result4-3.xlsx（紧急购电 {n_seg} 段）')

def run_stage_scan():
    print('\n4. scan: α-β 重扫（波动电价）:')
    rowsB = []
    gridB = [(1.00, 0.95), (1.00, 1.00), (1.10, 0.95), (1.10, 1.00),
             (1.22, 1.00), (1.35, 0.95), (1.35, 1.00)]
    for a, b in gridB:
        pkl = os.path.join(OUT, f'_q4_B_a{a}_b{b}.pkl')
        if os.path.exists(pkl):
            with open(pkl, 'rb') as f:
                st = pickle.load(f)
        else:
            t0 = time.time()
            st = q2_run('B', a, b, P4)
            with open(pkl, 'wb') as f:
                pickle.dump(st, f)
            print(f'      ({time.time()-t0:.0f}s)', flush=True)
        tot = st['plan_cost'].sum() + st['emerg_cost'].sum()
        rowsB.append((a, b, tot, st['emerg_energy'].sum()))
        print(f'   [Q2-B α={a:.2f},β={b:.2f}] {tot/1e4:,.1f} 万'
              f'（紧急 {st["emerg_energy"].sum()/1e4:.1f} 万kWh）', flush=True)
    pd.DataFrame(rowsB, columns=['alpha', 'beta', 'total', 'emerg_kWh']
                 ).to_csv(os.path.join(OUT, '_q4_scanB.csv'), index=False)

    rowsQ3 = []
    gridQ3 = [(1.00, 1.00), (1.22, 1.00), (1.35, 1.00)]
    for a, b in gridQ3:
        pkl = os.path.join(OUT, f'_q4_S1218_a{a}_b{b}.pkl')
        if os.path.exists(pkl):
            with open(pkl, 'rb') as f:
                st = pickle.load(f)
        else:
            t0 = time.time()
            st = q3_run((12, 18), a, b, P4)
            with open(pkl, 'wb') as f:
                pickle.dump(st, f)
            print(f'      ({time.time()-t0:.0f}s)', flush=True)
        rowsQ3.append((a, b, st['total'], st['em_energy']))
        print(f'   [Q3-S1218 α={a:.2f},β={b:.2f}] {st["total"]/1e4:,.1f} 万'
              f'（紧急 {st["em_energy"]/1e4:.1f} 万kWh）', flush=True)
    pd.DataFrame(rowsQ3, columns=['alpha', 'beta', 'total', 'emerg_kWh']
                 ).to_csv(os.path.join(OUT, '_q4_scanQ3.csv'), index=False)
    bestB = min(rowsB, key=lambda r: r[2])
    bestQ3 = min(rowsQ3, key=lambda r: r[2])
    print(f'\n   Q2-B 波动电价最优: α={bestB[0]:.2f}, β={bestB[1]:.2f} → '
          f'{bestB[2]/1e4:,.1f} 万')
    print(f'   Q3-S1218 波动电价最优: α={bestQ3[0]:.2f}, β={bestQ3[1]:.2f} → '
          f'{bestQ3[2]/1e4:,.1f} 万')

def run_stage_analyze(B=None, S=None):
    if B is None:
        with open(os.path.join(OUT, '_q4_B122.pkl'), 'rb') as f:
            B = pickle.load(f)
    if S is None:
        with open(os.path.join(OUT, '_q4_S1218.pkl'), 'rb') as f:
            S = pickle.load(f)
    print('\n5. analyze:')
    totB4 = B['plan_cost'].sum() + B['emerg_cost'].sum()
    print(f'   [锚点复现] A(1,1) P4 = {cand[ALIGN]/1e4:,.1f} 万'
          f'（教程 2844.1）；A(1.22) P4 见下；均值日计划按P4结算见下')
    stA122 = q2_run('A', 1.22, 1.00, P4)
    tA122 = stA122['plan_cost'].sum() + stA122['emerg_cost'].sum()
    stLed = q2m.run_year('A', q3m.p, q3m.Lm, q3m.Gm, q3m.L_all, q3m.G_all,
                         q3m.dates, q3m.mask_out, alpha=1.0, beta=1.0,
                         price_mat=np.tile(q3m.p, (365, 1)), ledger_mat=P4)
    tLed = stLed['plan_cost'].sum() + stLed['emerg_cost'].sum()
    print(f'   A(1.22) P4 = {tA122/1e4:,.1f} 万（教程锚点 2121.8）')
    print(f'   均值日计划按P4结算 = {tLed/1e4:,.1f} 万（教程锚点 2872.3）')
    print(f'   → 日前价格信息价值 ≈ {(tLed - cand[ALIGN])/1e4:,.1f} 万/年'
          f'（{tLed/1e4:,.1f} − {cand[ALIGN]/1e4:,.1f}；作为对照，'
          f'光伏预报信息价值 136 万/年）')
    print(f'\n   [固定 vs 波动]（交付口径 B(1.22,1.00) / Q3-S1218(1.22,1.00)）:')
    print(f'   Q2-B: 固定 1,979.4 万 → 波动 {totB4/1e4:,.1f} 万'
          f'（计划 {B["plan_cost"].sum()/1e4:,.1f}，紧急 '
          f'{B["emerg_cost"].sum()/1e4:,.1f}，占比 '
          f'{B["emerg_cost"].sum()/totB4*100:.1f}%）')
    print(f'   Q3-S1218: 固定 1,817.6 万 → 波动 {S["total"]/1e4:,.1f} 万'
          f'（计划 {S["plan"]/1e4:,.1f}，偏差 {S["dev"]/1e4:,.1f}，'
          f'紧急 {S["emerg"]/1e4:,.1f}）')

    i426 = int(np.where(q3m.dates == pd.Timestamp('2025-04-26'))[0][0])
    ds = '2025-04-26'
    d = B['detail'][ds]
    md = slice(60, 84)
    chg_426 = float(d['chg'][md].sum())
    all_mid = np.mean([B['detail'][str(dd.date())]['chg'][md].sum()
                       for dd in q3m.out_dates])
    pmin = float(P4[i426].min())
    print(f'\n   [近零日 4-26] 当日最低电价 {pmin:.4f} 元'
          f'（10:00-14:00 充电 {chg_426:,.0f} kWh vs 全年均值 '
          f'{all_mid:,.0f} kWh，{"高于" if chg_426 > all_mid else "不高于"}均值）')

    print('\n   [信息价值链·全文收尾]')
    print('   均值日计划(Q2·A基准 2692.2万)')
    print('     → +保守缓冲(P0, α=1.18/0.95)      省 610.6 万 → 2081.6')
    print('     → +执行自由化+系数迁移(B, 1.22/1.00) 省 102.2 万 → 1979.4')
    print('     → +0:00 光伏预报(Q3·S0)           省 135.9 万 → 1843.5')
    print('     → +12:00/18:00 调整(Q3·S1218)     省  25.9 万 → 1817.6')
    print(f'     → 换波动电价(Q4·B口径)             {totB4/1e4-1979.4:+,.1f} 万 '
          f'→ {totB4/1e4:,.1f}')
    print('     → 日前价格信息价值(对照实验)        ~28 万/年')

    fig, axes = plt.subplots(1, 2, figsize=(13, 5))
    ax = axes[0]
    labels = ['Q2-B 固定', 'Q2-B 波动', 'Q3-S1218 固定', 'Q3-S1218 波动']
    vals = [1979.4, totB4 / 1e4, 1817.6, S['total'] / 1e4]
    bars = ax.bar(labels, vals, color=['gray', 'steelblue', 'lightgray',
                                       'seagreen'], alpha=0.9)
    for bb, v in zip(bars, vals):
        ax.text(bb.get_x() + bb.get_width() / 2, v, f'{v:,.0f}',
                ha='center', va='bottom', fontsize=10)
    ax.set_ylabel('总费用 / 万元')
    ax.set_title('固定电价 vs 波动电价（交付口径）')
    ax.tick_params(axis='x', labelsize=9)
    ax = axes[1]
    i426 = i426
    xx = np.arange(1, T + 1) / 6
    ax.plot(xx, P4[i426], 'r-', lw=1.5, label='当日电价(左轴)')
    ax.set_ylabel('电价 / 元/kWh', color='r')
    ax.set_xlabel('时刻 / h')
    ax2 = ax.twinx()
    ax2.plot(xx, B['detail']['2025-04-26']['soc'], 'k-', lw=1.8,
             label='SOC(右轴)')
    ax2.set_ylabel('SOC / kWh', color='k')
    ax.set_title('2025-04-26 近零电价日：电价曲线 vs SOC')
    lines1, lb1 = ax.get_legend_handles_labels()
    lines2, lb2 = ax2.get_legend_handles_labels()
    ax.legend(lines1 + lines2, lb1 + lb2, fontsize=9)
    plt.tight_layout()
    fig.savefig(os.path.join(OUT, 'Q4分析.png'), dpi=150)
    print('\n图已保存: Q4分析.png')


if __name__ == '__main__':
    stage = sys.argv[1] if len(sys.argv) > 1 else 'all'
    if stage in ('final', 'all'):
        B4, S4 = run_stage_final()
    else:
        B4, S4 = None, None
    if stage in ('scan', 'all'):
        run_stage_scan()
    if stage in ('analyze', 'all'):
        run_stage_analyze(B4, S4)
    print('=' * 72)
    print('第四问完成')
    print('=' * 72)
