import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import pandas as pd
import numpy as np
import os
import sys
import time
sys.stdout.reconfigure(encoding='utf-8')

from scipy.optimize import linprog
from scipy.sparse import csc_matrix

plt.rcParams['font.family'] = 'SimHei'
plt.rcParams['axes.unicode_minus'] = False

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, '..', '附件')
OUTPUT_DIR = BASE_DIR

ETA = 0.9
SMIN, SMAX = 1200.0, 10800.0
S0 = 6000.0
PMAX = 5000.0
T = 144
DT = 1.0 / 6.0
EMAX = PMAX * DT

P0_REF_TOTAL = 26922467.55


def load_all():
    a1 = pd.read_excel(os.path.join(DATA_DIR, '附件1.xlsx'))
    load = pd.read_excel(os.path.join(DATA_DIR, '附件2.xlsx'), sheet_name='小区负载')
    pv = pd.read_excel(os.path.join(DATA_DIR, '附件2.xlsx'), sheet_name='光伏发电实际功率')
    return a1, load, pv


def solve_plan_lp(price, load, pv, dt=DT, eta=ETA, s0=S0,
                  smin=SMIN, smax=SMAX, pmax=PMAX, free_end=True):
    from scipy.sparse import csc_matrix
    Tn = len(price)
    Emax = pmax * dt
    n = 5 * Tn
    obj = np.concatenate([price, np.zeros(4 * Tn)])
    rows, cols, vals = [], [], []
    b_eq = np.zeros(2 * Tn)
    for t in range(Tn):
        rows += [t, t, t, t]
        cols += [t, 2 * Tn + t, Tn + t, 3 * Tn + t]
        vals += [1.0, 1.0, -1.0, -1.0]
        b_eq[t] = (load[t] - pv[t]) * dt
        r = Tn + t
        rows += [r, r, r]
        cols += [Tn + t, 2 * Tn + t, 4 * Tn + t]
        vals += [-eta, 1.0 / eta, 1.0]
        if t > 0:
            rows.append(r)
            cols.append(4 * Tn + t - 1)
            vals.append(-1.0)
    b_eq[Tn] = s0
    A_eq = csc_matrix((vals, (rows, cols)), shape=(2 * Tn, n))
    bounds = ([(0, None)] * Tn + [(0, Emax)] * Tn + [(0, Emax)] * Tn
              + [(0, float(u)) for u in pv * dt] + [(smin, smax)] * Tn)
    if free_end:
        bounds[5 * Tn - 1] = (smin, smax)
    else:
        bounds[5 * Tn - 1] = (s0, s0)
    res = linprog(obj, A_eq=A_eq, b_eq=b_eq, bounds=bounds, method='highs')
    assert res.status == 0, f'计划LP求解失败: {res.message}'
    x = res.x
    return x[:Tn], x[Tn:2 * Tn], x[2 * Tn:3 * Tn]


def build_templates(p, mode='cost'):
    FLAT = 5.0 * float(p.max())
    TILT = 0.25
    tpls = []
    for t0 in range(T):
        h = T - t0
        rows, cols, vals = [], [], []
        for j in range(h):
            rows += [j, j, j, j]
            cols += [h + j, 2 * h + j, j, 3 * h + j]
            vals += [1.0, 1.0, -1.0, -1.0]
        for j in range(h):
            rows += [h + j, h + j, h + j]
            cols += [4 * h + j, j, h + j]
            vals += [1.0, -ETA, 1.0 / ETA]
            if j > 0:
                rows.append(h + j)
                cols.append(4 * h + j - 1)
                vals.append(-1.0)
        A_eq = csc_matrix((vals, (rows, cols)), shape=(2 * h, 5 * h))
        obj = np.zeros(5 * h)
        if mode == 'cost':
            obj[2 * h:3 * h] = 5.0 * p[t0:t0 + h]
        else:
            obj[2 * h:3 * h] = FLAT + TILT * p[t0:t0 + h]
        obj[3 * h:4 * h] = 0.01
        bounds = ([(0.0, EMAX)] * h + [(0.0, EMAX)] * h
                  + [(0.0, None)] * h + [(0.0, None)] * h
                  + [(SMIN, SMAX)] * h)
        tpls.append((A_eq, obj, bounds, h))
    return tpls


def solve_step(tpl, x_vec, g_vec, l_vec, s, fix_idx=None):
    A_eq, obj, bounds, h = tpl
    b_eq = np.empty(2 * h)
    b_eq[:h] = l_vec - x_vec - g_vec
    b_eq[h] = s
    if h > 1:
        b_eq[h + 1:] = 0.0
    if fix_idx is not None:
        bounds = list(bounds)
        bounds[fix_idx] = (0.0, 0.0)
    res = linprog(obj, A_eq=A_eq, b_eq=b_eq, bounds=bounds, method='highs')
    return res


def execute_day_p0(x_p, c_p, d_p, L_act, G_act, e0):
    s = e0
    emerg = np.zeros(T); chg = np.zeros(T); dis = np.zeros(T)
    soc = np.zeros(T); cur = np.zeros(T)
    for t in range(T):
        d = min(d_p[t], max(0.0, (s - SMIN) * ETA))
        supply = x_p[t] + G_act[t] + d
        if supply >= L_act[t]:
            surplus, em = supply - L_act[t], 0.0
        else:
            surplus, em = 0.0, L_act[t] - supply
        c = min(c_p[t], surplus, max(0.0, (SMAX - s) / ETA))
        cur[t] = surplus - c
        s = s + ETA * c - d / ETA
        emerg[t], chg[t], dis[t], soc[t] = em, c, d, s
    return s, emerg, chg, dis, soc, cur


def execute_day_mpc(tpls, x_p, L_act, G_act, Lfut, Gfut, e0):
    s = e0
    n_branch = 0
    emerg = np.zeros(T); chg = np.zeros(T); dis = np.zeros(T)
    soc = np.zeros(T); cur = np.zeros(T)
    for t in range(T):
        h = T - t
        l_vec = np.empty(h)
        g_vec = np.empty(h)
        l_vec[0] = L_act[t]
        g_vec[0] = G_act[t]
        if h > 1:
            l_vec[1:] = Lfut[t + 1:]
            g_vec[1:] = Gfut[t + 1:]
        res = solve_step(tpls[t], x_p[t:], g_vec, l_vec, s)
        assert res.status == 0, f'执行LP失败 t={t}: {res.message}'
        sol = res.x
        c0, em0 = sol[0], sol[2 * h]
        if c0 > 1e-9 and em0 > 1e-9:
            n_branch += 1
            res_a = solve_step(tpls[t], x_p[t:], g_vec, l_vec, s, fix_idx=0)
            res_b = solve_step(tpls[t], x_p[t:], g_vec, l_vec, s,
                               fix_idx=2 * h)
            cands = [r for r in (res_a, res_b) if r.status == 0]
            best = min(cands, key=lambda r: r.fun)
            sol = best.x
            c0, em0 = sol[0], sol[2 * h]
        d0, cu0 = sol[h], sol[3 * h]
        s = s + ETA * c0 - d0 / ETA
        if SMIN - 1e-6 < s < SMIN:
            s = SMIN
        elif SMAX < s < SMAX + 1e-6:
            s = SMAX
        emerg[t], chg[t], dis[t], soc[t], cur[t] = em0, c0, d0, s, cu0
    return s, emerg, chg, dis, soc, cur, n_branch


def execute_day_clairvoyant(tpl, x_p, L_act, G_act, e0):
    res = solve_step(tpl, x_p, G_act, L_act, e0)
    assert res.status == 0, f'C口径LP失败: {res.message}'
    sol = res.x
    chg = sol[0:T]
    dis = sol[T:2 * T]
    emerg = sol[2 * T:3 * T]
    cur = sol[3 * T:4 * T]
    soc = sol[4 * T:5 * T]
    return float(soc[-1]), emerg, chg, dis, soc, cur


def run_year(mode, p, Lm, Gm, L_all, G_all, dates, mask_out,
             alpha=1.0, beta=1.0, collect=False, cache=None, tpls=None,
             tag='', price_mat=None, ledger_mat=None):
    if cache is None:
        cache = {}
    n_days = len(dates)
    if price_mat is None:
        price_mat = np.tile(p, (n_days, 1))
    if ledger_mat is None:
        ledger_mat = price_mat
    fixed_price = bool(np.allclose(price_mat, price_mat[0]))
    if fixed_price and tpls is None:
        tpls = build_templates(price_mat[0], 'cost')
    Lm_kwh = Lm * DT
    Gm_kwh = Gm * DT
    Lfut = Lm_kwh * alpha
    Gfut = Gm_kwh * beta
    e0 = S0
    n_plans = 0
    arb = 0
    n_branch = 0
    plan_cost, emerg_cost, emerg_energy = [], [], []
    monthly = np.zeros(12)
    detail = {}
    for i in range(n_days):
        p_day = price_mat[i]
        p_led = ledger_mat[i]
        e0_start = e0
        key = ((i,) if not fixed_price else ()) + \
              (round(alpha, 3), round(beta, 3), round(e0, 4))
        if key not in cache:
            cache[key] = solve_plan_lp(p_day, Lm * alpha, Gm * beta, s0=e0)
            n_plans += 1
        x_p, c_p, d_p = cache[key]
        if mode == 'A':
            e0, em, ch, di, so, cu = execute_day_p0(
                x_p, c_p, d_p, L_all[i], G_all[i], e0)
        elif mode == 'B':
            tpls_d = tpls if fixed_price else build_templates(p_day, 'cost')
            e0, em, ch, di, so, cu, nb = execute_day_mpc(
                tpls_d, x_p, L_all[i], G_all[i], Lfut, Gfut, e0)
            n_branch += nb
        else:
            tpl_d = tpls[0] if fixed_price else build_templates(p_day, 'cost')[0]
            e0, em, ch, di, so, cu = execute_day_clairvoyant(
                tpl_d, x_p, L_all[i], G_all[i], e0)
        monthly[dates.iloc[i].month - 1] += em.sum()
        arb += int(((ch > 1e-9) & (em > 1e-9)).sum())
        if mask_out[i]:
            plan_cost.append(float(p_led @ x_p))
            emerg_cost.append(float(5 * p_led @ em))
            emerg_energy.append(float(em.sum()))
            if collect:
                detail[str(dates.iloc[i].date())] = {
                    'x_p': x_p, 'chg': ch, 'dis': di, 'soc': so,
                    'emerg': em, 'cur': cu, 'e0': e0_start}
    st = {'plan_cost': np.array(plan_cost),
          'emerg_cost': np.array(emerg_cost),
          'emerg_energy': np.array(emerg_energy),
          'monthly': monthly, 'n_plans': n_plans, 'detail': detail,
          'arb': arb, 'n_branch': n_branch}
    if tag:
        tot = (st['plan_cost'].sum() + st['emerg_cost'].sum()) / 1e4
        print(f'   [{tag}] α={alpha:.2f} β={beta:.2f}: 总费用 {tot:,.1f} 万元'
              f'（紧急 {st["emerg_energy"].sum() / 1e4:.1f} 万kWh，'
              f'解计划LP {n_plans} 次）', flush=True)
    return st


def runs_of(mask):
    out, s = [], None
    for i, v in enumerate(mask):
        if v and s is None:
            s = i
        if not v and s is not None:
            out.append((s, i - 1))
            s = None
    if s is not None:
        out.append((s, len(mask) - 1))
    return out


def fmt_interval(j, k):
    def m2s(m):
        m = int(round(m))
        if m >= 1440:
            return '0:00+1'
        return f'{m // 60}:{m % 60:02d}'
    return f'{m2s(10 * j)}-{m2s(10 * (k + 1))}'


def merge_emergency(emerg, eps=1e-9):
    return [(fmt_interval(a, b), float(emerg[a:b + 1].sum()))
            for a, b in runs_of(emerg > eps)]


def main(alpha=1.0, beta=1.0):
    print('=' * 70)
    print(f'第二问·三口径：A·P0（对照）/ B·MPC-fair（主交付，α={alpha}, β={beta}）'
          f'/ C（诊断）')
    print('=' * 70)

    a1, load_df, pv_df = load_all()
    p = a1['电价'].values
    Lm = a1['小区负载'].values
    Gm = a1['光伏发电预测功率'].values
    dates = pd.to_datetime(load_df.iloc[:, 0])
    L_all = load_df.iloc[:, 1:].values / 6.0
    G_all = pv_df.iloc[:, 1:].values / 6.0
    mask_out = (dates >= pd.Timestamp('2025-02-01')).values
    idx_out = np.where(mask_out)[0]
    out_dates = dates.iloc[idx_out]
    print(f'\n1. 数据: {len(dates)}天, 正式统计 {int(mask_out.sum())} 天'
          f'（{out_dates.iloc[0].date()} ~ {out_dates.iloc[-1].date()}），1月预热')

    t0 = time.time()
    tpls_cost = build_templates(p, 'cost')
    tpls_energy = build_templates(p, 'energy')
    print(f'   执行层LP模板: 2×144 个（{time.time() - t0:.1f}s）')

    cache = {}

    print('\n2. A·P0（对照口径，自检）:')
    t0 = time.time()
    A = run_year('A', p, Lm, Gm, L_all, G_all, dates, mask_out,
                 collect=True, cache=cache, tpls=tpls_cost, tag='A·P0')
    tot_A = A['plan_cost'].sum() + A['emerg_cost'].sum()
    print(f'   计划购电费   = {A["plan_cost"].sum():,.2f} 元')
    print(f'   紧急购电费   = {A["emerg_cost"].sum():,.2f} 元')
    print(f'   总费用       = {tot_A:,.2f} 元（参考 {P0_REF_TOTAL:,.2f}，'
          f'差 {tot_A - P0_REF_TOTAL:+,.2f}）')
    ok_ref = abs(tot_A - P0_REF_TOTAL) < 1.0
    print(f'   自检 {"✓" if ok_ref else "✗ 偏差>1元，需排查!"}'
          f'（耗时 {time.time() - t0:.1f}s）')

    print('\n3. C口径（仅诊断，非策略）:')
    t0 = time.time()
    C1 = run_year('C', p, Lm, Gm, L_all, G_all, dates, mask_out,
                  alpha=alpha, beta=beta,
                  cache=cache, tpls=tpls_cost, tag='C1·绝对下界')
    tot_C1 = C1['plan_cost'].sum() + C1['emerg_cost'].sum()
    C2 = run_year('C', p, Lm, Gm, L_all, G_all, dates, mask_out,
                  alpha=alpha, beta=beta,
                  cache=cache, tpls=tpls_energy, tag='C2·能量地板')
    tot_C2 = C2['plan_cost'].sum() + C2['emerg_cost'].sum()
    print(f'   C1 = {tot_C1:,.2f} 元（实际价+套利自由度全放开 → 任何执行的费用下界；'
          f'含套利时段 {C1["arb"]} 个）')
    print(f'   C2 = {tot_C2:,.2f} 元（最小紧急电量解的费用结算，仅参考；'
          f'紧急电量即合法能量地板，套利时段 {C2["arb"]} 个）')
    print(f'   （耗时 {time.time() - t0:.1f}s）')

    print('\n4. B·MPC-fair（主交付口径，约4分钟）:')
    t0 = time.time()
    B = run_year('B', p, Lm, Gm, L_all, G_all, dates, mask_out,
                 alpha=alpha, beta=beta,
                 collect=True, cache=cache, tpls=tpls_cost, tag='B·MPC')
    tot_B = B['plan_cost'].sum() + B['emerg_cost'].sum()
    print(f'   计划购电费   = {B["plan_cost"].sum():,.2f} 元')
    print(f'   紧急购电费   = {B["emerg_cost"].sum():,.2f} 元')
    print(f'   总费用       = {tot_B:,.2f} 元')
    print(f'   紧急购电总量 = {B["emerg_energy"].sum() / 1e4:.1f} 万 kWh')
    print(f'   发生紧急购电天数 = {int((B["emerg_energy"] > 0).sum())} / 334')
    print(f'   （耗时 {time.time() - t0:.1f}s，反套利分支触发 {B["n_branch"]} 次）')

    print('\n5. 汇总（2-12月，334天）:')
    print(f'   {"口径":<16}{"计划费(万)":>12}{"紧急费(万)":>12}'
          f'{"总费用(万)":>12}{"紧急量(万kWh)":>14}')
    for nm, st, tot in [('A·P0 对照', A, tot_A), ('B·MPC 主交付', B, tot_B),
                        ('C1·绝对下界', C1, tot_C1), ('C2·能量地板', C2, tot_C2)]:
        print(f'   {nm:<16}{st["plan_cost"].sum() / 1e4:>12,.1f}'
              f'{st["emerg_cost"].sum() / 1e4:>12,.1f}{tot / 1e4:>12,.1f}'
              f'{st["emerg_energy"].sum() / 1e4:>14,.1f}')
    ok_order = tot_C1 <= tot_B + 1e-6 and tot_B <= tot_A + 1e-6
    print(f'   费用序 C1 ≤ B ≤ A: {"✓" if ok_order else "✗"}')
    A_em, B_em, C2_em = (A['emerg_energy'].sum(), B['emerg_energy'].sum(),
                         C2['emerg_energy'].sum())
    print(f'   紧急量分解（全实测，万kWh）: P0 {A_em/1e4:.1f} = '
          f'执行自由化消除 {((A_em-B_em)/1e4):.1f} + 信息与因果代价 '
          f'{((B_em-C2_em)/1e4):.1f} + 合法能量地板 {C2_em/1e4:.1f}')
    print(f'   ★ 1500万检验: 绝对下界 {tot_C1/1e4:,.0f} 万 >> 1500 万'
          f' → 执行层无法到达 1500 万，"其他队1500万"只能来自计划层口径差异')

    print('\n6. 指定日期（B口径，result2 同源）:')
    print(f'   {"日期":>10}{"计划购电kWh":>13}{"计划费元":>12}'
          f'{"紧急kWh":>10}{"紧急费元":>12}')
    for ds in ['2025-03-20', '2025-06-21', '2025-09-23', '2025-12-21']:
        d = B['detail'][ds]
        print(f'   {ds:>10}{d["x_p"].sum():>13.2f}{float(p @ d["x_p"]):>12.2f}'
              f'{d["emerg"].sum():>10.1f}{float(5 * p @ d["emerg"]):>12.2f}')
    print('   对照（A口径同四日紧急购电量 kWh）:')
    for ds in ['2025-03-20', '2025-06-21', '2025-09-23', '2025-12-21']:
        print(f'   {ds}: A={A["detail"][ds]["emerg"].sum():>10.1f}  '
              f'B={B["detail"][ds]["emerg"].sum():>10.1f}')

    print('\n7. 验证清单:')
    all_soc = np.concatenate([B['detail'][str(d.date())]['soc'] for d in out_dates])
    ok_soc = all_soc.min() >= SMIN - 1e-4 and all_soc.max() <= SMAX + 1e-4
    print(f'   B口径SOC范围 = [{all_soc.min():.2f}, {all_soc.max():.2f}] '
          f'{"✓" if ok_soc else "✗"}')
    ds_v = '2025-12-21'
    i_v = int(np.where(dates == pd.Timestamp(ds_v))[0][0])
    d_v = B['detail'][ds_v]
    lhs = d_v['x_p'].sum() + G_all[i_v].sum() + d_v['dis'].sum() + d_v['emerg'].sum()
    rhs = L_all[i_v].sum() + d_v['chg'].sum() + d_v['cur'].sum()
    ok_bal = abs(lhs - rhs) < 1e-6
    print(f'   能量守恒(12-21,B): {lhs:.6f} vs {rhs:.6f} '
          f'{"✓" if ok_bal else "✗"}')
    plan_sums = np.array([B['detail'][str(d.date())]['x_p'].sum() for d in out_dates])
    print(f'   B口径日计划购电: min {plan_sums.min():.2f} / max {plan_sums.max():.2f} kWh'
          f'（e0漂移→计划逐日微调，非不动点）')
    print(f'   紧急电充电套利时段（c·em同时>0）: B={B["arb"]}（分支强制，触发 '
          f'{B["n_branch"]} 次）/ C1={C1["arb"]}（下界允许）/ C2={C2["arb"]}（应≈0）')

    print('\n8. 填写 result2.xlsx（B口径）:')
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(DATA_DIR, '附件5', 'result2.xlsx'))
    ws1 = wb['计划购电量']
    for j in range(len(idx_out)):
        ds = str(out_dates.iloc[j].date())
        x_p = B['detail'][ds]['x_p']
        roll = np.roll(x_p, -1)
        r = 2 + j
        for k in range(T):
            ws1.cell(row=r, column=2 + k, value=round(float(roll[k]), 2))
        ws1.cell(row=r, column=146, value=round(float(x_p.sum()), 2))
        ws1.cell(row=r, column=147, value=round(float(p @ x_p), 2))
    ws2 = wb['充放电量']
    if ws2.max_row > 1:
        ws2.delete_rows(2, ws2.max_row - 1)
    blocks = ['0:00-4:00', '4:00-8:00', '8:00-12:00',
              '12:00-16:00', '16:00-20:00', '20:00-24:00']
    r = 2
    for j in range(len(idx_out)):
        ds = str(out_dates.iloc[j].date())
        d = B['detail'][ds]
        for b in range(6):
            sl = slice(b * 24, (b + 1) * 24)
            ws2.cell(row=r, column=1,
                     value=out_dates.iloc[j] if b == 0 else None)
            ws2.cell(row=r, column=2, value=blocks[b])
            ws2.cell(row=r, column=3, value=round(float(d['chg'][sl].sum()), 2))
            ws2.cell(row=r, column=4, value=round(float(d['dis'][sl].sum()), 2))
            if b == 0:
                ws2.cell(row=r, column=5, value='0:00')
                ws2.cell(row=r, column=6, value=round(float(d['e0']), 2))
            elif b == 1:
                ws2.cell(row=r, column=5, value='24:00')
                ws2.cell(row=r, column=6, value=round(float(d['soc'][-1]), 2))
            r += 1
    ws3 = wb['紧急购电量']
    if ws3.max_row > 1:
        ws3.delete_rows(2, ws3.max_row - 1)
    r = 2
    n_seg = 0
    for j in range(len(idx_out)):
        ds = str(out_dates.iloc[j].date())
        ivs = merge_emergency(B['detail'][ds]['emerg'])
        for m_idx, (label, energy) in enumerate(ivs):
            ws3.cell(row=r, column=1,
                     value=out_dates.iloc[j] if m_idx == 0 else None)
            ws3.cell(row=r, column=2, value=label)
            ws3.cell(row=r, column=3, value=round(energy, 2))
            r += 1
            n_seg += 1
    out_path = os.path.join(OUTPUT_DIR, 'result2.xlsx')
    wb.save(out_path)
    print(f'   已保存 {out_path}（紧急购电 {n_seg} 段）')

    months = np.arange(1, 13)
    fig, ax = plt.subplots(figsize=(10, 4.5))
    w = 0.27
    ax.bar(months - w, A['monthly'] / 1e4, width=w, label='A·P0 对照',
           color='gray', alpha=0.85)
    ax.bar(months, B['monthly'] / 1e4, width=w, label='B·MPC 主交付',
           color='steelblue', alpha=0.85)
    ax.bar(months + w, C2['monthly'] / 1e4, width=w,
           label='C2·能量地板', color='seagreen', alpha=0.85)
    ax.set_xlabel('月份')
    ax.set_ylabel('紧急购电量 / 万kWh')
    ax.set_title('月度紧急购电分布：三口径对比（1月为预热段）')
    ax.set_xticks(months)
    ax.legend()
    plt.tight_layout()
    fig.savefig(os.path.join(OUTPUT_DIR, 'Q2月度紧急购电.png'), dpi=150)

    fig, ax = plt.subplots(figsize=(8.5, 4.5))
    names = ['A·P0\n对照', 'B·MPC\n主交付', 'C1·绝对下界\n(含套利自由度)',
             'C2·能量地板\n(最小紧急电量)']
    tots = [tot_A / 1e4, tot_B / 1e4, tot_C1 / 1e4, tot_C2 / 1e4]
    colors = ['gray', 'steelblue', 'indianred', 'seagreen']
    bars = ax.bar(names, tots, color=colors, alpha=0.85)
    for bb, v in zip(bars, tots):
        ax.text(bb.get_x() + bb.get_width() / 2, v, f'{v:,.0f}',
                ha='center', va='bottom', fontsize=11)
    ax.set_ylabel('总费用 / 万元')
    ax.set_title('四口径全年总费用对比（2025.2.1–12.31）')
    ax.set_ylim(0, max(tots) * 1.15)
    plt.tight_layout()
    fig.savefig(os.path.join(OUTPUT_DIR, 'Q2三口径对比.png'), dpi=150)

    fig, axes = plt.subplots(2, 2, figsize=(13, 7))
    for col, (ds, name) in enumerate([('2025-12-21', '12-21 冬至'),
                                      ('2025-06-21', '6-21 夏至')]):
        d = B['detail'][ds]
        x = np.arange(1, T + 1) / 6
        axes[0, col].plot(x, d['soc'], 'k-', lw=1.8)
        axes[0, col].axhline(SMIN, color='gray', ls=':', lw=1)
        axes[0, col].axhline(SMAX, color='gray', ls='--', lw=1)
        axes[0, col].set_ylabel('SOC / kWh')
        axes[0, col].set_title(f'{name}：B口径 SOC 轨迹')
        axes[0, col].grid(alpha=0.3)
        axes[1, col].bar(x, d['emerg'], width=1 / 6, color='red', alpha=0.7)
        axes[1, col].set_ylabel('紧急购电 / kWh/10min')
        axes[1, col].set_xlabel('时刻 / h')
        axes[1, col].set_title(f'紧急购电量（合计 {d["emerg"].sum():.0f} kWh）')
        axes[1, col].grid(alpha=0.3)
    plt.tight_layout()
    fig.savefig(os.path.join(OUTPUT_DIR, 'Q2案例对比.png'), dpi=150)
    print('   图已保存: Q2月度紧急购电.png / Q2三口径对比.png / Q2案例对比.png')

    print('=' * 70)
    print('主阶段完成。α-β 扫描请运行: python q2_main.py scan')
    print('=' * 70)


def scan():
    print('=' * 70)
    print('α-β 扫描（B·MPC-fair 口径，12 组合）')
    print('=' * 70)
    a1, load_df, pv_df = load_all()
    p = a1['电价'].values
    Lm = a1['小区负载'].values
    Gm = a1['光伏发电预测功率'].values
    dates = pd.to_datetime(load_df.iloc[:, 0])
    L_all = load_df.iloc[:, 1:].values / 6.0
    G_all = pv_df.iloc[:, 1:].values / 6.0
    mask_out = (dates >= pd.Timestamp('2025-02-01')).values
    tpls = build_templates(p)
    cache = {}
    combos = [(a, b) for a in (1.00, 1.05, 1.10, 1.15)
              for b in (0.90, 0.95, 1.00)]
    rows = []
    csv_path = os.path.join(OUTPUT_DIR, 'Q2_scan_results.csv')
    for a, b in combos:
        t0 = time.time()
        st = run_year('B', p, Lm, Gm, L_all, G_all, dates, mask_out,
                      alpha=a, beta=b, cache=cache, tpls=tpls,
                      tag=f'α={a:.2f},β={b:.2f}')
        tot = st['plan_cost'].sum() + st['emerg_cost'].sum()
        rows.append((a, b, tot, st['emerg_energy'].sum()))
        pd.DataFrame(rows, columns=['alpha', 'beta', 'total', 'emerg_kwh']
                     ).to_csv(csv_path, index=False)
        print(f'      ({time.time() - t0:.0f}s，已存盘 {csv_path})', flush=True)
    best = min(rows, key=lambda r: r[2])
    print(f'\n   最优组合: α={best[0]:.2f}, β={best[1]:.2f} → '
          f'{best[2]:,.0f} 元')
    base = [r for r in rows if r[0] == 1.00 and r[1] == 1.00][0]
    print(f'   基准 (1,1): {base[2]:,.0f} 元；'
          f'保守计划红利 {(base[2] - best[2]) / 1e4:,.1f} 万元')

    fig, ax = plt.subplots(figsize=(7, 5.5))
    aa = np.array([r[0] for r in rows])
    bb = np.array([r[1] for r in rows])
    tt = np.array([r[2] for r in rows]) / 1e4
    sc = ax.scatter(aa, bb, c=tt, cmap='viridis_r', s=140, edgecolor='k')
    plt.colorbar(sc, ax=ax, label='总费用 / 万元')
    ax.scatter([best[0]], [best[1]], c='red', marker='*', s=300, zorder=5,
               label=f'最优 α={best[0]:.2f}, β={best[1]:.2f}')
    ax.scatter([1.0], [1.0], c='red', marker='X', s=140, label='基准 (1,1)')
    for a, b, v in [(r[0], r[1], r[2] / 1e4) for r in rows]:
        ax.annotate(f'{v:,.0f}', (a, b), textcoords='offset points',
                    xytext=(6, 6), fontsize=8)
    ax.set_xlabel('α（计划负载高估系数）')
    ax.set_ylabel('β（计划光伏折扣系数）')
    ax.set_title('B口径（MPC-fair）下的保守计划灵敏度')
    ax.legend()
    plt.tight_layout()
    fig.savefig(os.path.join(OUTPUT_DIR, 'Q2_alpha_beta扫描.png'), dpi=150)
    print('   图已保存: Q2_alpha_beta扫描.png')
    print('=' * 70)
    print('扫描完成')
    print('=' * 70)


if __name__ == '__main__':
    if len(sys.argv) > 1 and sys.argv[1] == 'scan':
        scan()
    elif len(sys.argv) > 1 and sys.argv[1] == 'fill':
        main(float(sys.argv[2]), float(sys.argv[3]))
    else:
        main()
