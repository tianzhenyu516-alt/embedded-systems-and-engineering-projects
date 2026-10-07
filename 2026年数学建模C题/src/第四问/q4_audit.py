import numpy as np
import pandas as pd
import pickle
import os
import sys
import openpyxl
sys.stdout.reconfigure(encoding='utf-8')

BASE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(BASE, '..', '第三问'))
import q3_main as q3m
import q2_main as q2m

p = q3m.p
P4 = q3m.a4.iloc[:, 1:].values.astype(float)
T = 144
TOL = 0.0051
TOL_FEE = 0.011
idx_out = q3m.idx_out
out_dates = q3m.out_dates
issues_total = []


def audit_q2like(xlsx_path, pkl_path, price, seg_expect, label, has_adj):
    with open(pkl_path, 'rb') as f:
        det = pickle.load(f)['detail']
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    issues = []
    ws = wb[wb.sheetnames[0]]
    field = 'xhat' if 'xhat' in det[str(out_dates.iloc[0].date())] else 'x_p'
    bad = 0
    for j in range(334):
        i = idx_out[j]
        ds = str(out_dates.iloc[j].date())
        x = det[ds][field]
        roll = np.roll(x, -1)
        for k in range(T):
            v = ws.cell(row=2 + j, column=2 + k).value
            if v is None or abs(float(v) - round(float(roll[k]), 2)) > TOL:
                bad += 1
                if bad <= 3:
                    issues.append(f'{label}/{ws.title} r{2+j}c{2+k}: {v} ≠ {round(float(roll[k]),2)}')
        v146 = ws.cell(row=2 + j, column=146).value
        if abs(float(v146) - round(float(x.sum()), 2)) > TOL_FEE:
            issues.append(f'{label}/{ws.title} r{2+j} c146: {v146} ≠ {x.sum():.2f}')
        fee = float(price @ x) if np.ndim(price) == 1 else float(price[i] @ x)
        v147 = ws.cell(row=2 + j, column=147).value
        if abs(float(v147) - round(fee, 2)) > TOL_FEE:
            issues.append(f'{label}/{ws.title} r{2+j} c147: {v147} ≠ {fee:.2f}')
    print(f'   [{label}] {ws.title}: 334行×146格 {"全部一致 ✓" if bad == 0 else f"{bad} 格不符"}')

    if has_adj:
        ws = wb[wb.sheetnames[1]]
        field = 'x_final'
        bad = 0
        for j in range(334):
            i = idx_out[j]
            ds = str(out_dates.iloc[j].date())
            x = det[ds][field]
            roll = np.roll(x, -1)
            for k in range(T):
                v = ws.cell(row=2 + j, column=2 + k).value
                if v is None or abs(float(v) - round(float(roll[k]), 2)) > TOL:
                    bad += 1
                    if bad <= 3:
                        issues.append(f'{label}/{ws.title} r{2+j}c{2+k}: {v} ≠ {round(float(roll[k]),2)}')
            v146 = ws.cell(row=2 + j, column=146).value
            if abs(float(v146) - round(float(x.sum()), 2)) > TOL_FEE:
                issues.append(f'{label}/{ws.title} r{2+j} c146: {v146} ≠ {x.sum():.2f}')
            fee = float(price @ x) if np.ndim(price) == 1 else float(price[i] @ x)
            v147 = ws.cell(row=2 + j, column=147).value
            if abs(float(v147) - round(fee, 2)) > TOL_FEE:
                issues.append(f'{label}/{ws.title} r{2+j} c147: {v147} ≠ {fee:.2f}')
        print(f'   [{label}] {ws.title}: 334行×146格 {"全部一致 ✓" if bad == 0 else f"{bad} 格不符"}')

    ws_i = 2 if has_adj else 1
    ws = wb[wb.sheetnames[ws_i]]
    bad = 0
    blocks_ok = True
    for j in range(334):
        ds = str(out_dates.iloc[j].date())
        d = det[ds]
        for b in range(6):
            r = 2 + j * 6 + b
            sl = slice(b * 24, (b + 1) * 24)
            v3 = ws.cell(row=r, column=3).value
            v4 = ws.cell(row=r, column=4).value
            if abs(float(v3) - round(float(d['chg'][sl].sum()), 2)) > TOL:
                bad += 1
                if bad <= 3:
                    issues.append(f'{label}/充放 r{r}c3: {v3} ≠ {d["chg"][sl].sum():.2f}')
            if abs(float(v4) - round(float(d['dis'][sl].sum()), 2)) > TOL:
                bad += 1
                if bad <= 3:
                    issues.append(f'{label}/充放 r{r}c4: {v4} ≠ {d["dis"][sl].sum():.2f}')
            if b == 0:
                v6 = ws.cell(row=r, column=6).value
                if abs(float(v6) - round(float(d['e0']), 2)) > TOL:
                    bad += 1
                    issues.append(f'{label}/充放 r{r} 0:00储电量: {v6} ≠ {d["e0"]:.2f}')
            elif b == 1:
                v6 = ws.cell(row=r, column=6).value
                if abs(float(v6) - round(float(d['soc'][-1]), 2)) > TOL:
                    bad += 1
                    issues.append(f'{label}/充放 r{r} 24:00储电量: {v6} ≠ {d["soc"][-1]:.2f}')
    chain_bad = 0
    for j in range(333):
        d0 = det[str(out_dates.iloc[j].date())]
        d1 = det[str(out_dates.iloc[j + 1].date())]
        if abs(d1['e0'] - d0['soc'][-1]) > 1e-6:
            chain_bad += 1
            if chain_bad <= 3:
                issues.append(f'{label}/SOC链 {out_dates.iloc[j+1].date()}: '
                              f'e0 {d1["e0"]:.4f} ≠ 前日末SOC {d0["soc"][-1]:.4f}')
    print(f'   [{label}] {ws.title}: 充放电格 {"全部一致 ✓" if bad == 0 else f"{bad} 格不符"}；'
          f'SOC链 e₀(t+1)=24:00(t): {"✓" if chain_bad == 0 else f"{chain_bad} 处断裂"}')
    issues.extend([] if bad == 0 and chain_bad == 0 else [])

    ws_i = 3 if has_adj else 2
    ws = wb[wb.sheetnames[ws_i]]
    n_seg = 0
    em_sum = 0.0
    r = 2
    while True:
        v_label = ws.cell(row=r, column=2).value
        v_val = ws.cell(row=r, column=3).value
        if v_label is None and v_val is None:
            break
        n_seg += 1
        em_sum += float(v_val)
        r += 1
    em_total = sum(float(det[str(d.date())]['emerg'].sum()) for d in out_dates)
    ok_seg = (n_seg == seg_expect)
    ok_sum = abs(em_sum - em_total) < n_seg * 0.005 + 0.05
    print(f'   [{label}] {ws.title}: 紧急段数 {n_seg}（预期 {seg_expect}）'
          f'{"✓" if ok_seg else "✗"}；段电量合计 {em_sum:,.1f} vs pickle '
          f'{em_total:,.1f} {"✓" if ok_sum else "✗"}')
    if not ok_seg:
        issues.append(f'{label}/紧急段数 {n_seg} ≠ {seg_expect}')
    if not ok_sum:
        issues.append(f'{label}/紧急电量合计 {em_sum} ≠ {em_total}')
    return issues


def check_426():
    with open(os.path.join(BASE, '_q4_B122.pkl'), 'rb') as f:
        B = pickle.load(f)
    d = B['detail']['2025-04-26']
    v = float(d['chg'][60:84].sum())
    ok = abs(v - 20000.0) < 1.0
    print(f'   [result4-2] 4-26 午间(10:00-14:00)充电 = {v:,.1f} kWh '
          f'{"✓（功率上限满充）" if ok else "✗"}')
    if not ok:
        issues_total.append(f'4-26 午间充电 {v} ≠ 20000')


print('=' * 72)
print('Q3/Q4 交付文件逐格对账（只读）')
print('=' * 72)

issues_total += audit_q2like(
    os.path.join(BASE, '..', '第三问', 'result3.xlsx'),
    os.path.join(BASE, '..', '第三问', '_q3v2_S1218_a1.22.pkl'),
    p, 577, 'result3', has_adj=True)

issues_total += audit_q2like(
    os.path.join(BASE, 'result4-2.xlsx'),
    os.path.join(BASE, '_q4_B122.pkl'),
    P4, 943, 'result4-2', has_adj=False)

issues_total += audit_q2like(
    os.path.join(BASE, 'result4-3.xlsx'),
    os.path.join(BASE, '_q4_S1218.pkl'),
    P4, 634, 'result4-3', has_adj=True)

check_426()

print('-' * 72)
if not issues_total:
    print('对账结论：三份交付文件全部逐格一致 ✓✓✓（最后一道锁落下）')
else:
    print(f'对账结论：发现 {len(issues_total)} 处不符：')
    for s in issues_total[:30]:
        print('   -', s)


def audit_result2():
    pkl2 = os.path.join(BASE, '..', '第二问', '_q2_B122.pkl')
    if not os.path.exists(pkl2):
        st2 = q2m.run_year('B', q3m.p, q3m.Lm,
                           q3m.a1['光伏发电预测功率'].values, q3m.L_all,
                           q3m.G_all, q3m.dates, q3m.mask_out, alpha=1.22,
                           beta=1.00, collect=True)
        with open(pkl2, 'wb') as f:
            pickle.dump(st2, f)
        print(f'   [Q2-B 重跑] 总 {st2["plan_cost"].sum()+st2["emerg_cost"].sum():,.2f} '
              f'（vs 封版 19,793,691.58）')
    iss2 = audit_q2like(
        os.path.join(BASE, '..', '第二问', 'result2.xlsx'),
        pkl2, q3m.p, 917, 'result2', has_adj=False)
    print('-' * 72)
    if not iss2:
        print('result2 对账结论：全部逐格一致 ✓')
    else:
        print(f'result2 对账结论：{len(iss2)} 处不符')
        for s in iss2[:20]:
            print('   -', s)


if len(sys.argv) > 1 and sys.argv[1] == 'q2':
    audit_result2()
print('=' * 72)
